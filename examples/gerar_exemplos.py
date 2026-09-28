"""
Gera dados de exemplo SINTETICOS para os quatro protocolos do BSP.

Nenhum destes ficheiros corresponde a uma pessoa real. As trajectorias do
centro de pressao (CoP) sao somas de sinusoides com ruido aleatorio, com
semente fixa para que o resultado seja sempre igual.

Uso (a partir da pasta examples/):
    python gerar_exemplos.py
"""

import math
import os
import random

from openpyxl import Workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
DATA = "01-01-2026"


def trajectoria(n, passo_ms, amp, rng, amp_extra=None):
    """Devolve lista de (t_ms, x, y). amp_extra(t_ms) permite aumentar a
    oscilacao em partes do ensaio (ex.: fora da janela de analise)."""
    fx1, fx2 = rng.uniform(0.2, 0.5), rng.uniform(0.8, 1.5)
    fy1, fy2 = rng.uniform(0.2, 0.5), rng.uniform(0.8, 1.5)
    px, py = rng.uniform(0, 6.28), rng.uniform(0, 6.28)
    dx = dy = 0.0
    pts = []
    for i in range(n):
        t = i * passo_ms
        s = t / 1000.0
        a = amp * (amp_extra(t) if amp_extra else 1.0)
        dx += rng.gauss(0, 0.05 * a)
        dy += rng.gauss(0, 0.05 * a)
        dx *= 0.98
        dy *= 0.98
        x = a * (0.6 * math.sin(2 * math.pi * fx1 * s + px) + 0.3 * math.sin(2 * math.pi * fx2 * s)) + dx
        y = a * (0.6 * math.cos(2 * math.pi * fy1 * s + py) + 0.3 * math.sin(2 * math.pi * fy2 * s)) + dy
        pts.append((t, x, y))
    return pts


def escrever_plataforma(caminho, medicao, pts, forca=700.0):
    """Formato de exportacao da plataforma (texto separado por tabs, extensao .xls)."""
    with open(caminho, "w", encoding="iso-8859-1", newline="\r\n") as f:
        f.write(f"Stability export for measurement: {medicao}\n")
        f.write("Patient name: Exemplo sintetico\n")
        f.write(f"Measurement done on {DATA}\n")
        f.write("\n")
        f.write("Frame\tTime (ms)\tCOF X (mm)\tCOF Y (mm)\tForce (N)\tFrameDist (mm)\n")
        xa = ya = None
        for i, (t, x, y) in enumerate(pts):
            d = 0.0 if xa is None else math.hypot(x - xa, y - ya)
            xa, ya = x, y
            f.write(f"{i}\t{t:.0f}\t{x:.3f}\t{y:.3f}\t{forca:.1f}\t{d:.3f}\n")


def escrever_arco(caminho, medicao, pts, forca=650.0):
    """Formato Tiro com Arco (colunas 'Entire plate COF', 50 Hz)."""
    with open(caminho, "w", encoding="iso-8859-1", newline="\r\n") as f:
        f.write(f"Stability export for measurement: {medicao}\n")
        f.write("Patient name: Exemplo sintetico\n")
        f.write(f"Measurement done on {DATA}\n")
        f.write("\n")
        f.write("Frame\tTime (ms)\tEntire plate COF X (mm)\tEntire plate COF Y (mm)\t"
                "Entire plate COF Force (N)\n")
        for i, (t, x, y) in enumerate(pts):
            f.write(f"{i}\t{t:.0f}\t{x:.3f}\t{y:.3f}\t{forca:.1f}\n")


def gerar_fms_unipodal(pasta_base, rng, amp_base):
    """FMS Bipodal e Apoio Unipodal: dir_1..5 e esq_1..5 por individuo."""
    os.makedirs(pasta_base, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "inicio_fim"
    ws.append(["Atleta"] + [f"{k}{i}" for i in range(1, 11) for k in ("ini", "fim")])
    for n, nome in enumerate(["Exemplo_A", "Exemplo_B", "Exemplo_C"], start=1):
        pasta = os.path.join(pasta_base, f"{n:02d}_{nome}")
        os.makedirs(pasta, exist_ok=True)
        linha = [nome]
        for lado, off in (("dir", 0), ("esq", 5)):
            amp = amp_base * rng.uniform(0.8, 1.3)
            for t in range(1, 6):
                pts = trajectoria(2000, 10.0, amp, rng)   # 20 s a 100 Hz
                escrever_plataforma(os.path.join(pasta, f"{lado}_{t} - {DATA} - Stability export.xls"),
                                    f"{lado}_{t}", pts)
        for _ in range(10):
            ini = rng.choice([2000, 2500, 3000])
            linha += [ini, ini + 15000]
        ws.append(linha)
    wb.save(os.path.join(pasta_base, "inicio_fim.xlsx"))


def gerar_tiro(pasta_base, rng):
    """Tiro: trial{dist}_{ensaio} por individuo e ficheiro de tempos com 3 folhas."""
    os.makedirs(pasta_base, exist_ok=True)
    dist, n_ens = "5m", 5
    wb = Workbook()
    folhas = {
        "toque": wb.active,
        "pontaria": wb.create_sheet(),
        "disparo": wb.create_sheet(),
    }
    folhas["toque"].title = "tempo (toque)"
    folhas["pontaria"].title = "tempo (pontaria)"
    folhas["disparo"].title = "tempo (disparo)"
    cab = ["Individuo"] + [f"{dist} ({i})" for i in range(1, n_ens + 1)]
    for chave, ws in folhas.items():
        ws.append([f"Tempos de {chave} em ms: colunas B-F camara, colunas G-K placa"])
        ws.append(cab + [f"{dist} ({i})" for i in range(1, n_ens + 1)])

    for ind in (1, 2, 3):
        pasta = os.path.join(pasta_base, f"{ind}_Exemplo_{'ABC'[ind - 1]}")
        os.makedirs(pasta, exist_ok=True)
        cam = {"toque": [], "pontaria": [], "disparo": []}
        placa = []
        for t in range(1, n_ens + 1):
            toque = rng.randint(800, 1200)
            pont = toque + rng.randint(1800, 2500)
            disp = pont + rng.randint(2500, 3500)

            def amp_extra(tm, toque=toque, pont=pont, disp=disp):
                if tm < pont:
                    return 3.0          # levantar a arma
                if tm < disp:
                    return 1.0          # pontaria
                if tm < disp + 300:
                    return 2.5          # recuo do disparo
                return 1.5

            pts = trajectoria(1000, 10.0, 2.0 * rng.uniform(0.8, 1.3), rng, amp_extra)  # 10 s
            escrever_plataforma(os.path.join(pasta, f"trial5_{t} - 01_01_2026 - Stability_export.xls"),
                                f"trial5_{t}", pts)
            desvio_camara = rng.randint(200, 400)
            cam["toque"].append(toque + desvio_camara)
            cam["pontaria"].append(pont + desvio_camara)
            cam["disparo"].append(disp + desvio_camara)
            placa.append(toque)
        folhas["toque"].append([ind] + cam["toque"] + placa)
        folhas["pontaria"].append([ind] + cam["pontaria"])
        folhas["disparo"].append([ind] + cam["disparo"])
    wb.save(os.path.join(pasta_base, "tempos_tiro.xlsx"))


def gerar_arco(pasta_base, rng, n_ens=30):
    """Tiro com Arco: {id}_{ensaio} por individuo e Inicio_fim_vfinal.xlsx."""
    os.makedirs(pasta_base, exist_ok=True)
    wb = Workbook()
    ws_t = wb.active
    ws_t.title = "tempo do toque"
    ws_1 = wb.create_sheet("confirmação_1")
    ws_2 = wb.create_sheet("confirmação_2")
    cab = ["atleta_id"] + [f"ensaio_{i}" for i in range(1, n_ens + 1)]
    for ws in (ws_t, ws_1, ws_2):
        ws.append(cab)

    for pid, nome in ((901, "Exemplo_A"), (902, "Exemplo_B")):
        pasta = os.path.join(pasta_base, f"{pid}_{nome}")
        os.makedirs(pasta, exist_ok=True)
        lt, l1, l2 = [pid], [pid], [pid]
        for t in range(1, n_ens + 1):
            toque = rng.randint(400, 800)
            c1 = toque + rng.randint(1500, 2200)
            c2 = c1 + rng.randint(1500, 2500)

            def amp_extra(tm, c1=c1, c2=c2):
                return 1.0 if c1 <= tm <= c2 else 4.0

            pts = trajectoria(300, 20.0, 1.5 * rng.uniform(0.8, 1.3), rng, amp_extra)  # 6 s a 50 Hz
            escrever_arco(os.path.join(pasta, f"{pid}_{t} - {DATA} - Stability export.xls"),
                          f"{pid}_{t}", pts)
            lt.append(toque)
            l1.append(c1)
            l2.append(c2)
        ws_t.append(lt)
        ws_1.append(l1)
        ws_2.append(l2)
    wb.save(os.path.join(pasta_base, "Inicio_fim_vfinal.xlsx"))


def main():
    rng = random.Random(2026)
    gerar_fms_unipodal(os.path.join(AQUI, "fms_bipodal"), rng, amp_base=3.0)
    gerar_fms_unipodal(os.path.join(AQUI, "unipodal"), rng, amp_base=6.0)
    gerar_tiro(os.path.join(AQUI, "tiro"), rng)
    gerar_arco(os.path.join(AQUI, "tiro_arco"), rng)
    print("Dados de exemplo gerados em", AQUI)


if __name__ == "__main__":
    main()
